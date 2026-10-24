"""Assign categories (rules, then your manual choices) and detect transfers."""
import re
from dataclasses import dataclass
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from .models import TRANSFER, UNCATEGORIZED, is_yes

OVERRIDE_COLUMNS = ["txn_id", "category", "is_transfer", "date", "account", "description", "amount"]

# Column headings in the output workbook that users may fill in
COL_ID = "ID"
COL_YOUR_CATEGORY = "Your category"
COL_YOUR_TRANSFER = "Transfer? (Yes/No)"
COL_MAKE_RULE = "Make a rule from this text"

TRANSFER_HINTS = re.compile(r"PAYMENT|TRANSFER|\bTFR\b|PYMT|PMT|PREAUTH|BILL PAY")


@dataclass
class Rule:
    keyword: str
    category: str
    is_transfer: bool
    applies_to: str  # "any" | "money in" | "money out"

    def __post_init__(self):
        self._rx = re.compile(r"(?<![A-Z0-9])" + re.escape(self.keyword.upper()))

    def matches(self, desc_upper: str, amount: float) -> bool:
        if self.applies_to == "money in" and amount <= 0:
            return False
        if self.applies_to == "money out" and amount >= 0:
            return False
        return bool(self._rx.search(desc_upper))


def is_income_category(category: str) -> bool:
    return str(category).strip().lower().startswith("income")


# ---------------------------------------------------------------- loading

def load_rules(categories_path: Path) -> list[Rule]:
    ws = load_workbook(categories_path, read_only=True, data_only=True)["Rules"]
    rules = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        kw, cat, transfer, applies = (list(row) + [None] * 4)[:4]
        if not kw or not str(kw).strip():
            continue
        is_t = is_yes(transfer)
        rules.append(Rule(
            keyword=str(kw).strip(),
            category=str(cat).strip() if cat else (TRANSFER if is_t else UNCATEGORIZED),
            is_transfer=is_t,
            applies_to=str(applies or "any").strip().lower(),
        ))
    return rules


def load_budget(categories_path: Path) -> dict[str, float]:
    wb = load_workbook(categories_path, read_only=True, data_only=True)
    if "Budget" not in wb.sheetnames:
        return {}
    budget = {}
    for row in wb["Budget"].iter_rows(min_row=2, values_only=True):
        cat, amt = (list(row) + [None, None])[:2]
        if cat and isinstance(amt, (int, float)):
            budget[str(cat).strip()] = float(amt)
    return budget


def load_overrides(path: Path) -> pd.DataFrame:
    if not path.exists():
        return pd.DataFrame(columns=OVERRIDE_COLUMNS)
    return pd.read_csv(path, dtype=str, keep_default_na=False)


# ---------------------------------------------------------------- harvesting edits

def harvest_edits(output_xlsx: Path, overrides_path: Path, categories_path: Path) -> tuple[int, int]:
    """Copy what the user typed into the last report into overrides.csv / Categories.xlsx.

    Returns (overrides_saved, rules_added).
    """
    if not output_xlsx.exists():
        return 0, 0
    wb = load_workbook(output_xlsx, read_only=True, data_only=True)
    edits: dict[str, dict] = {}
    new_rules: list[tuple[str, str, str]] = []
    for sheet in ("Review", "Transactions"):
        if sheet not in wb.sheetnames:
            continue
        rows = wb[sheet].iter_rows(values_only=True)
        header = next(rows, None)
        if not header or COL_ID not in header:
            continue
        idx = {h: i for i, h in enumerate(header) if h}

        def get(row, col):
            i = idx.get(col)
            v = row[i] if i is not None and i < len(row) else None
            return str(v).strip() if v is not None else ""

        for row in rows:
            txn_id = get(row, COL_ID)
            cat, transfer, rule_text = get(row, COL_YOUR_CATEGORY), get(row, COL_YOUR_TRANSFER), get(row, COL_MAKE_RULE)
            if not txn_id or not (cat or transfer):
                continue
            is_t = "yes" if is_yes(transfer) else ("no" if transfer else "")
            edits[txn_id] = {
                "txn_id": txn_id, "category": cat, "is_transfer": is_t,
                "date": get(row, "Date")[:10], "account": get(row, "Account"),
                "description": get(row, "Description"), "amount": get(row, "Amount"),
            }
            if rule_text and (cat or is_t == "yes"):
                new_rules.append((rule_text, cat or TRANSFER, "Yes" if is_t == "yes" else ""))
    wb.close()

    if edits:
        existing = load_overrides(overrides_path)
        existing = existing[~existing["txn_id"].isin(edits.keys())]
        merged = pd.concat([existing, pd.DataFrame(edits.values())], ignore_index=True)
        merged[OVERRIDE_COLUMNS].to_csv(overrides_path, index=False)

    added = _add_rules(categories_path, new_rules) if new_rules else 0
    return len(edits), added


def _add_rules(categories_path: Path, new_rules: list[tuple[str, str, str]]) -> int:
    """Insert user-made rules at the top so they win over the starter rules."""
    wb = load_workbook(categories_path)
    ws = wb["Rules"]
    have = {
        (str(r[0]).strip().upper(), str(r[1]).strip())
        for r in ws.iter_rows(min_row=2, values_only=True) if r and r[0]
    }
    added = 0
    for kw, cat, transfer in new_rules:
        if (kw.upper(), cat) in have:
            continue
        ws.insert_rows(2)
        for col, val in enumerate([kw, cat, transfer, "Any"], start=1):
            ws.cell(row=2, column=col, value=val)
        have.add((kw.upper(), cat))
        added += 1
    if added:
        wb.save(categories_path)
    return added


# ---------------------------------------------------------------- applying

def categorize(ledger: pd.DataFrame, rules: list[Rule], overrides: pd.DataFrame) -> pd.DataFrame:
    df = ledger.copy()
    cats, transfers, sources = [], [], []
    for desc, amt in zip(df["description"].str.upper(), df["amount"]):
        rule = next((r for r in rules if r.matches(desc, amt)), None)
        if rule:
            cats.append(rule.category); transfers.append(rule.is_transfer); sources.append("Rule")
        else:
            cats.append(UNCATEGORIZED); transfers.append(False); sources.append("")
    df["category"], df["is_transfer"], df["source"] = cats, transfers, sources

    for o in overrides.to_dict("records"):
        mask = df["txn_id"] == o["txn_id"]
        if not mask.any():
            continue
        if o.get("is_transfer") == "yes":
            df.loc[mask, "is_transfer"] = True
            df.loc[mask, "category"] = o["category"] or TRANSFER
        elif o.get("is_transfer") == "no":
            df.loc[mask, "is_transfer"] = False
            if o["category"]:
                df.loc[mask, "category"] = o["category"]
            elif (df.loc[mask, "category"] == TRANSFER).all():
                df.loc[mask, "category"] = UNCATEGORIZED
        elif o["category"]:
            df.loc[mask, "category"] = o["category"]
            df.loc[mask, "is_transfer"] = o["category"].strip().lower() == TRANSFER.lower()
        df.loc[mask, "source"] = "You"
    return detect_transfers(df)


def detect_transfers(df: pd.DataFrame, window_days: int = 3) -> pd.DataFrame:
    """Pair up equal-and-opposite amounts between two of your accounts within a few days.

    If either side looks like a payment/transfer the pair is marked as a transfer
    automatically; otherwise both are flagged for review as a possible transfer.
    Choices you made yourself (source == "You") are never changed.
    """
    df = df.copy()
    df["review_note"] = ""
    open_rows = df[(df["source"] != "You") & (df["amount"] != 0)]
    used: set[int] = set()
    for i, row in open_rows[open_rows["amount"] < 0].iterrows():
        cands = open_rows[
            (open_rows["amount"] == -row["amount"])
            & (open_rows["account"] != row["account"])
            & ((open_rows["date"] - row["date"]).abs().dt.days <= window_days)
            & (~open_rows.index.isin(used))
        ]
        if cands.empty:
            continue
        j = (cands["date"] - row["date"]).abs().idxmin()
        pair = [i, j]
        used.update(pair)
        if df.loc[pair, "is_transfer"].all():
            continue
        hinted = any(TRANSFER_HINTS.search(str(d).upper()) for d in df.loc[pair, "description"])
        if hinted or df.loc[pair, "is_transfer"].any():
            for k in pair:
                if not df.at[k, "is_transfer"]:
                    df.at[k, "is_transfer"] = True
                    df.at[k, "category"] = TRANSFER
                    df.at[k, "source"] = "Auto (matched transfer)"
        else:
            other = {i: j, j: i}
            for k in pair:
                o = df.loc[other[k]]
                df.at[k, "review_note"] = (
                    f"Possible transfer: matches {o['amount']:+,.2f} in {o['account']} "
                    f"on {o['date']:%Y-%m-%d}")
    return df
