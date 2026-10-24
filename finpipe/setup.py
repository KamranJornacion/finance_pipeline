"""First-run setup: create the user-editable workbooks and the statements/ folders.

These files are generated here (rather than committed) so personal edits never
end up in git and a fresh clone always starts from clean defaults.
"""
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .models import ACCOUNT_TYPES, BANKS, Account, is_yes

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")

ACCOUNTS_HEADERS = ["Folder name", "Bank", "Account type", "Display name", "Flip signs?"]
EXAMPLE_ACCOUNTS = [
    ("td-chequing", "TD", "Chequing", "TD Chequing", ""),
    ("tangerine-savings", "Tangerine", "Savings", "Tangerine Savings", ""),
    ("amex-card", "Amex", "Credit card", "Amex Card", ""),
    ("cibc-visa", "CIBC", "Credit card", "CIBC Visa", ""),
]

RULES_HEADERS = ["If description contains", "Category", "Is transfer?", "Applies to"]
# Order matters: the first matching row wins. Transfers come first so card payments
# are never mistaken for spending or income.
STARTER_RULES = [
    # transfers between your own accounts / credit card payments
    ("PAYMENT THANK YOU", "Transfer", "Yes", "Any"),
    ("PAYMENT - THANK YOU", "Transfer", "Yes", "Any"),
    ("PAYMENT RECEIVED", "Transfer", "Yes", "Any"),
    ("PAIEMENT MERCI", "Transfer", "Yes", "Any"),
    ("TFR-TO", "Transfer", "Yes", "Any"),
    ("TFR-FR", "Transfer", "Yes", "Any"),
    ("TRANSFER TO", "Transfer", "Yes", "Any"),
    ("TRANSFER FROM", "Transfer", "Yes", "Any"),
    ("INTERNET TRANSFER", "Transfer", "Yes", "Any"),
    ("TD VISA PREAUTH", "Transfer", "Yes", "Any"),
    ("CIBC VISA", "Transfer", "Yes", "Money out"),
    ("AMEX BILL", "Transfer", "Yes", "Money out"),
    ("AMERICAN EXPRESS", "Transfer", "Yes", "Money out"),
    # income
    ("PAYROLL", "Income", "", "Money in"),
    ("DIRECT DEP", "Income", "", "Money in"),
    ("DIR DEP", "Income", "", "Money in"),
    ("PAY/PAY", "Income", "", "Money in"),
    ("INTEREST", "Income: Interest", "", "Money in"),
    # fees
    ("INTEREST", "Fees & Interest", "", "Money out"),
    ("SERVICE CHARGE", "Fees & Interest", "", "Any"),
    ("MONTHLY FEE", "Fees & Interest", "", "Any"),
    ("ANNUAL FEE", "Fees & Interest", "", "Any"),
    ("NSF", "Fees & Interest", "", "Any"),
    ("OVERDRAFT", "Fees & Interest", "", "Any"),
    # housing & bills
    ("MORTGAGE", "Housing", "", "Money out"),
    ("HYDRO", "Utilities", "", "Any"),
    ("ENBRIDGE", "Utilities", "", "Any"),
    ("ALECTRA", "Utilities", "", "Any"),
    ("ROGERS", "Phone & Internet", "", "Any"),
    ("BELL CANADA", "Phone & Internet", "", "Any"),
    ("TELUS", "Phone & Internet", "", "Any"),
    ("FIDO", "Phone & Internet", "", "Any"),
    ("KOODO", "Phone & Internet", "", "Any"),
    ("FREEDOM MOBILE", "Phone & Internet", "", "Any"),
    ("VIRGIN PLUS", "Phone & Internet", "", "Any"),
    ("INSURANCE", "Insurance", "", "Any"),
    # groceries
    ("LOBLAWS", "Groceries", "", "Any"),
    ("NO FRILLS", "Groceries", "", "Any"),
    ("SUPERSTORE", "Groceries", "", "Any"),
    ("METROLINX", "Transportation", "", "Any"),
    ("METRO", "Groceries", "", "Any"),
    ("SOBEYS", "Groceries", "", "Any"),
    ("FRESHCO", "Groceries", "", "Any"),
    ("FOOD BASICS", "Groceries", "", "Any"),
    ("ZEHRS", "Groceries", "", "Any"),
    ("FORTINOS", "Groceries", "", "Any"),
    ("FARM BOY", "Groceries", "", "Any"),
    ("T&T", "Groceries", "", "Any"),
    ("WHOLE FOODS", "Groceries", "", "Any"),
    ("SAVE ON FOODS", "Groceries", "", "Any"),
    ("COSTCO", "Groceries", "", "Any"),
    # dining
    ("UBER EATS", "Dining & Coffee", "", "Any"),
    ("UBEREATS", "Dining & Coffee", "", "Any"),
    ("DOORDASH", "Dining & Coffee", "", "Any"),
    ("SKIPTHEDISHES", "Dining & Coffee", "", "Any"),
    ("TIM HORTONS", "Dining & Coffee", "", "Any"),
    ("STARBUCKS", "Dining & Coffee", "", "Any"),
    ("MCDONALD", "Dining & Coffee", "", "Any"),
    ("RESTAURANT", "Dining & Coffee", "", "Any"),
    ("PIZZA", "Dining & Coffee", "", "Any"),
    # transportation
    ("PRESTO", "Transportation", "", "Any"),
    ("UBER", "Transportation", "", "Any"),
    ("LYFT", "Transportation", "", "Any"),
    ("GO TRANSIT", "Transportation", "", "Any"),
    ("PETRO-CANADA", "Transportation", "", "Any"),
    ("ESSO", "Transportation", "", "Any"),
    ("SHELL", "Transportation", "", "Any"),
    ("PIONEER", "Transportation", "", "Any"),
    ("PARKING", "Transportation", "", "Any"),
    ("IMPARK", "Transportation", "", "Any"),
    # subscriptions
    ("NETFLIX", "Subscriptions", "", "Any"),
    ("SPOTIFY", "Subscriptions", "", "Any"),
    ("DISNEY PLUS", "Subscriptions", "", "Any"),
    ("CRAVE", "Subscriptions", "", "Any"),
    ("APPLE.COM/BILL", "Subscriptions", "", "Any"),
    ("YOUTUBE", "Subscriptions", "", "Any"),
    ("PRIME VIDEO", "Subscriptions", "", "Any"),
    # health
    ("SHOPPERS DRUG", "Health & Pharmacy", "", "Any"),
    ("REXALL", "Health & Pharmacy", "", "Any"),
    ("PHARMACY", "Health & Pharmacy", "", "Any"),
    # shopping
    ("AMAZON", "Shopping", "", "Any"),
    ("AMZN", "Shopping", "", "Any"),
    ("WALMART", "Shopping", "", "Any"),
    ("CANADIAN TIRE", "Shopping", "", "Any"),
    ("BEST BUY", "Shopping", "", "Any"),
    ("IKEA", "Shopping", "", "Any"),
    ("WINNERS", "Shopping", "", "Any"),
    # cash
    ("ATM", "Cash", "", "Money out"),
]

BUDGET_HEADERS = ["Category", "Monthly budget"]
STARTER_BUDGET = [
    ("Housing", 2000), ("Utilities", 150), ("Phone & Internet", 150), ("Insurance", 150),
    ("Groceries", 600), ("Dining & Coffee", 250), ("Transportation", 250),
    ("Subscriptions", 60), ("Health & Pharmacy", 75), ("Shopping", 300),
    ("Cash", 100), ("Fees & Interest", 0),
]

HOW_TO_DOWNLOAD = """\
HOW TO DOWNLOAD YOUR STATEMENTS
================================

1. Sign in to online banking on a computer (not the app).
2. Open the account, look for "Download" / "Export transactions".
3. Choose the CSV / "Spreadsheet" / "Comma separated" format and a date range.
   Overlapping date ranges are fine - duplicates are removed automatically.
4. Save (or move) the file into the matching folder here, e.g. td-chequing/.
   Each folder is one account; folders are listed in Accounts.xlsx.
5. Double-click "Update Finances".

Per-bank hints (menu names change from time to time):
  TD        - EasyWeb > account > "Download" > Spreadsheet (.csv)
  CIBC      - account > "Download transactions" > Spreadsheet (.csv)
  Tangerine - account > "Download transactions" > Comma Separated Values (.csv)
  Amex      - "Statements & Activity" > "Download" > CSV

Adding a new account: add a row to Accounts.xlsx, run "Update Finances" once
and its folder will be created here.
"""


def _style_header(ws, headers, widths):
    ws.append(headers)
    for i, w in enumerate(widths, start=1):
        c = ws.cell(row=1, column=i)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        ws.column_dimensions[c.column_letter].width = w
    ws.freeze_panes = "A2"


def _add_dropdown(ws, col_letter, options, max_row=500):
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{col_letter}2:{col_letter}{max_row}")


def _instructions(ws, lines):
    ws.column_dimensions["A"].width = 110
    for line in lines:
        ws.append([line])
    ws["A1"].font = Font(bold=True, size=14)
    for row in ws.iter_rows():
        row[0].alignment = Alignment(wrap_text=True, vertical="top")


def create_accounts_workbook(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Accounts"
    _style_header(ws, ACCOUNTS_HEADERS, [24, 14, 16, 26, 14])
    for row in EXAMPLE_ACCOUNTS:
        ws.append(list(row))
    _add_dropdown(ws, "B", BANKS)
    _add_dropdown(ws, "C", ACCOUNT_TYPES)
    _add_dropdown(ws, "E", ["Yes"])
    _instructions(wb.create_sheet("How to use"), [
        "Your accounts",
        "One row per bank account or credit card you want to track.",
        "Folder name: a short name with no spaces, e.g. td-chequing. A folder with this name is "
        "created inside 'statements' - save that account's downloads there.",
        "Bank / Account type: pick from the drop-down.",
        "Display name: how the account appears in the report.",
        "Flip signs?: leave blank. Only set to Yes if this account's spending shows up as income "
        "in the report.",
        "Delete the example rows you don't need.",
    ])
    wb.save(path)


def create_categories_workbook(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Rules"
    _style_header(ws, RULES_HEADERS, [34, 24, 14, 14])
    for row in STARTER_RULES:
        ws.append(list(row))
    _add_dropdown(ws, "C", ["Yes"], 2000)
    _add_dropdown(ws, "D", ["Any", "Money in", "Money out"], 2000)

    bs = wb.create_sheet("Budget")
    _style_header(bs, BUDGET_HEADERS, [24, 16])
    for row in STARTER_BUDGET:
        bs.append(list(row))
    for r in range(2, 200):
        bs.cell(row=r, column=2).number_format = '"$"#,##0'

    _instructions(wb.create_sheet("How to use"), [
        "Categories and budget",
        "RULES: each row says 'if a transaction's description contains this text, give it this "
        "category'. The text must match the start of a word (so NSF won't match TRANSFER). "
        "Capital letters don't matter. Rules are checked top to bottom and the first "
        "match wins, so put more specific text (UBER EATS) above general text (UBER).",
        "Is transfer?: Yes for money moving between your own accounts (e.g. paying your credit "
        "card). Transfers are left out of income and spending so nothing is counted twice.",
        "Applies to: Any, or only Money in / Money out (e.g. INTEREST earned vs. charged).",
        "Categories whose name starts with 'Income' count as income; everything else is spending.",
        "BUDGET: your monthly target for each spending category.",
        "Tip: you rarely need to edit this file. In Finances.xlsx, type a category in the "
        "'Your category' column and (optionally) a word in 'Make a rule from this text' - "
        "the next update saves it for you.",
    ])
    wb.save(path)


def ensure_setup(root: Path) -> list[str]:
    """Create whatever is missing. Returns plain-English notes about what was created."""
    notes = []
    accounts_path = root / "Accounts.xlsx"
    categories_path = root / "Categories.xlsx"
    statements = root / "statements"
    if not accounts_path.exists():
        create_accounts_workbook(accounts_path)
        notes.append("Created Accounts.xlsx - open it and list your accounts.")
    if not categories_path.exists():
        create_categories_workbook(categories_path)
        notes.append("Created Categories.xlsx with starter rules and a sample budget.")
    statements.mkdir(exist_ok=True)
    howto = statements / "HOW_TO_DOWNLOAD.txt"
    if not howto.exists():
        howto.write_text(HOW_TO_DOWNLOAD)
    for acct in load_accounts(accounts_path):
        folder = statements / acct.folder
        if not folder.exists():
            folder.mkdir(parents=True)
            notes.append(f"Created folder statements/{acct.folder}/ for {acct.name}.")
    return notes


def load_accounts(path: Path) -> list[Account]:
    ws = load_workbook(path, read_only=True, data_only=True)["Accounts"]
    accounts = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        folder, bank, acct_type, name, flip = (list(row) + [None] * 5)[:5]
        if not folder or not str(folder).strip():
            continue
        folder = str(folder).strip()
        bank_match = next((b for b in BANKS if b.lower() == str(bank or "").strip().lower()), None)
        accounts.append(Account(
            folder=folder,
            name=str(name).strip() if name else folder,
            bank=bank_match or str(bank or "").strip(),
            type=str(acct_type or "").strip(),
            flip_signs=is_yes(flip),
        ))
    return accounts
