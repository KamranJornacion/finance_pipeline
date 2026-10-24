from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from finpipe.categorize import Rule, categorize, load_overrides
from finpipe.cli import main
from finpipe.report import cashflow_frame, monthly_cashflow


def _ledger(rows):
    df = pd.DataFrame(rows, columns=["txn_id", "date", "account", "description", "amount"])
    df["date"] = pd.to_datetime(df["date"])
    return df


def test_rules_first_match_wins_word_start_and_direction():
    rules = [Rule("UBER EATS", "Dining", False, "any"), Rule("UBER", "Transport", False, "any"),
             Rule("NSF", "Fees", False, "any"), Rule("INTEREST", "Income: Interest", False, "money in"),
             Rule("INTEREST", "Fees", False, "money out")]
    df = categorize(_ledger([
        ("a", "2026-01-01", "X", "UBER EATS TORONTO", -10),
        ("b", "2026-01-02", "X", "UBER *TRIP", -12),
        ("c", "2026-01-03", "X", "INTERAC E-TRANSFER", -50),   # NSF must not match inside TRANSFER
        ("d", "2026-01-04", "X", "INTEREST", 2),
        ("e", "2026-01-05", "X", "INTEREST CHARGE", -3),
    ]), rules, load_overrides(Path("/nonexistent")))
    assert df["category"].tolist() == ["Dining", "Transport", "Uncategorized", "Income: Interest", "Fees"]


def test_override_beats_rule():
    rules = [Rule("AMAZON", "Shopping", False, "any")]
    ov = pd.DataFrame([{"txn_id": "a", "category": "Gifts", "is_transfer": ""}])
    df = categorize(_ledger([("a", "2026-01-01", "X", "AMAZON.CA", -30)]), rules, ov)
    assert (df.loc[0, "category"], df.loc[0, "source"]) == ("Gifts", "You")


def test_transfer_pairing():
    df = categorize(_ledger([
        ("a", "2026-01-01", "Chequing", "ONLINE PAYMENT VISA", -400),
        ("b", "2026-01-02", "Visa", "THANKS", 400),
        ("c", "2026-01-05", "Chequing", "MISC", -99),
        ("d", "2026-01-06", "Savings", "DEPOSIT", 99),
    ]), [], load_overrides(Path("/nonexistent")))
    assert df["is_transfer"].tolist() == [True, True, False, False]   # hinted pair auto-matched
    assert df.loc[2, "review_note"].startswith("Possible transfer")   # unhinted pair goes to Review


def test_end_to_end_cashflow_and_review_roundtrip(project):
    assert main(["--root", str(project), "--no-open"]) == 0
    tx = pd.read_csv(project / "output/transactions.csv")
    flows = tx[~tx["is_transfer"]]
    monthly = monthly_cashflow(cashflow_frame(tx.assign(date=pd.to_datetime(tx["date"]))))
    assert round(monthly["Net"].sum(), 2) == round(flows["amount"].sum(), 2)
    sep = monthly.set_index("Month").loc["2026-09"]
    assert round(sep["Income"], 2) == 6401.23 and round(sep["Spending"], 2) == 2101.74

    # user fills in the Review sheet in Excel...
    xlsx = project / "output/Finances.xlsx"
    wb = load_workbook(xlsx)
    ws = wb["Review"]
    header = [c.value for c in ws[1]]
    col = {h: i + 1 for i, h in enumerate(header)}
    for r in range(2, ws.max_row + 1):
        desc = ws.cell(r, col["Description"]).value
        if desc == "MYSTERY SHOP":
            ws.cell(r, col["Your category"], "Shopping")
            ws.cell(r, col["Make a rule from this text"], "MYSTERY")
        elif desc.startswith("SEND E-TFR"):
            ws.cell(r, col["Your category"], "Housing")
        elif desc in ("MISC 0045", "EFT Deposit"):
            ws.cell(r, col["Transfer? (Yes/No)"], "Yes")
    wb.save(xlsx)

    # ...and runs the update again
    assert main(["--root", str(project), "--no-open"]) == 0
    tx = pd.read_csv(project / "output/transactions.csv").set_index("description")
    assert tx.loc["MYSTERY SHOP", "category"] == "Shopping"
    assert tx.loc["SEND E-TFR ***F3k", "category"] == "Housing"
    assert bool(tx.loc["MISC 0045", "is_transfer"]) and bool(tx.loc["EFT Deposit", "is_transfer"])
    assert pd.read_excel(project / "output/Finances.xlsx", "Review").empty
    rules = load_workbook(project / "Categories.xlsx")["Rules"]
    assert [c.value for c in rules[2]][:2] == ["MYSTERY", "Shopping"]   # new rule added at the top

    # edits survive further runs (the yellow columns are blank now, overrides.csv remembers)
    assert main(["--root", str(project), "--no-open"]) == 0
    tx = pd.read_csv(project / "output/transactions.csv").set_index("description")
    assert tx.loc["SEND E-TFR ***F3k", "category"] == "Housing"
